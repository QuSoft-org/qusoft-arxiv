import os
import json
import arxiv
import ssl
import certifi
import feedparser
from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from unicodedata import normalize
from datetime import datetime

CACHE_FILE = "query_cache.json"

def normalize_name(name):
    return normalize('NFKD', name).encode('ascii', 'ignore').decode('ascii').lower()


def get_qusoft_people(file_path):
    workbook = load_workbook(file_path)
    sheet = workbook.active
    qusoft_people = [
        normalize_name(row[0].value) for row in sheet.iter_rows(min_row=3, max_col=1) if row[0].value
    ]
    return qusoft_people


def load_cache():
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, "r") as f:
            return json.load(f)
    return {}


def save_cache(cache):
    with open(CACHE_FILE, "w") as f:
        json.dump(cache, f)


def search_arxiv_for_authors(qusoft_people, cache):
    client = arxiv.Client()
    results = []
    for person in qusoft_people:
        cache_key = f"arxiv:{person}"
        if cache_key in cache:
            person_results = cache[cache_key]
        else:
            print(f"Searching arXiv for papers by {person}...")
            search = arxiv.Search(
                query=f'au:"{person}"',
                max_results=10,
                sort_by=arxiv.SortCriterion.SubmittedDate
            )
            person_results = []
            try:
                for result in client.results(search):
                    highlighted_authors = []
                    qusoft_author_count = 0
                    for author in result.authors:
                        normalized_author = normalize_name(author.name)
                        if normalized_author in qusoft_people:
                            highlighted_authors.append(f"**{author.name}**")
                            qusoft_author_count += 1
                        else:
                            highlighted_authors.append(author.name)
                    entry = {
                        "source": "arXiv",
                        "id": result.entry_id.split("/")[-1],
                        "link": result.entry_id,
                        "authors": ", ".join(highlighted_authors),
                        "title": result.title,
                        "revision": str(result.updated.replace(tzinfo=None)),
                        "qusoft_author_count": qusoft_author_count
                    }
                    person_results.append(entry)
            except arxiv.UnexpectedEmptyPageError as e:
                print(f"Warning: No results for {person} ({e})")
            cache[cache_key] = person_results
        results.extend(person_results)
    return results


def search_iacr_for_authors(qusoft_people, cache):
    cache_key = "iacr"
    if cache_key in cache:
        return cache[cache_key]
    rss_url = "https://eprint.iacr.org/rss/rss.xml"
    ssl_context = ssl._create_unverified_context()
    feed = feedparser.parse(rss_url, request_headers={})
    results = []
    for entry in feed.entries:
        authors = entry.get("authors", [])
        highlighted_authors = []
        qusoft_author_count = 0
        for author in authors:
            normalized_author = normalize_name(author)
            if normalized_author in qusoft_people:
                highlighted_authors.append(f"**{author}**")
                qusoft_author_count += 1
            else:
                highlighted_authors.append(author)
        if qusoft_author_count > 0:
            pub_date = datetime(*entry.published_parsed[:6])
            result_entry = {
                "source": "IACR ePrint",
                "id": entry.link.split("/")[-1],
                "link": entry.link,
                "authors": ", ".join(highlighted_authors),
                "title": entry.title,
                "revision": str(pub_date),
                "qusoft_author_count": qusoft_author_count
            }
            results.append(result_entry)
    cache[cache_key] = results
    return results


def save_to_excel(data, output_file):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QuSoft Papers"
    headers = ["source", "id", "list of authors", "title", "revision", "qusoft_author_count"]
    sheet.append(headers)
    for entry in data:
        row = [
            entry["source"],
            entry["id"],
            entry["authors"],
            entry["title"],
            entry["revision"],
            entry["qusoft_author_count"]
        ]
        sheet.append(row)
        id_cell = sheet.cell(row=sheet.max_row, column=2)
        id_cell.hyperlink = entry["link"]
        id_cell.style = "Hyperlink"
    table_range = f"A1:F{sheet.max_row}"
    table = Table(displayName="QuSoftTable", ref=table_range)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=True
    )
    table.tableStyleInfo = style
    sheet.add_table(table)
    workbook.save(output_file)
    print(f"Data saved to {output_file}")


def main():
    qusoft_people_file = "24-04-22-QuSoftDB.xlsx"
    qusoft_people = get_qusoft_people(qusoft_people_file)
    cache = load_cache()
    arxiv_papers = search_arxiv_for_authors(qusoft_people, cache)
    # iacr_papers = search_iacr_for_authors(qusoft_people, cache)
    all_papers = arxiv_papers
    output_file = "qusoft_papers.xlsx"
    save_to_excel(all_papers, output_file)
    save_cache(cache)

if __name__ == "__main__":
    main()
