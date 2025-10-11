#!/usr/bin/env python3
"""
QuSoft arXiv Paper Scanner and PDF Generator

This script combines functionality to:
1. Search arXiv for papers by QuSoft authors using advanced Unicode-aware search strategies
2. Generate an Excel spreadsheet with all found papers
3. Create annotated PDF first pages with author portraits and QR codes

The script uses intelligent search strategies to handle Unicode characters in author names
and provides comprehensive caching to avoid redundant API calls.

Author: QuSoft Team
Date: October 2025
"""

import arxiv
import os
import json
import requests
import tempfile
from datetime import datetime
from io import BytesIO
from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import qrcode
import logging
import urllib3
import unicodedata

# ============================================================================
# CONFIGURATION
# ============================================================================

# Directory structure
CACHE_DIR = "cache"
CACHE_FILE = os.path.join(CACHE_DIR, "query_cache.json")
PORTRAITS_DIR = "qusofters/portraits"
PDF_DOWNLOAD_DIR = os.path.join("cache", "downloaded_pdfs")
TMP_DIR = "tmp_firstpages"
QUSOFT_PEOPLE_FILE = "24-04-22-QuSoftDB.xlsx"

# Output files
EXCEL_OUTPUT = "qusoft_papers.xlsx"
PDF_OUTPUT = "qusoft_firstpages_2025_Oct.pdf"

# Filter configuration
STARTS_WITH = "2510"  # Only include papers starting with this (e.g., "25" for 2025 papers)

# Feature flags
SUBMISSION_TO_QIP_BOX = False  # Add QIP submission checkbox to PDFs
ENABLE_HTTP_LOGGING = False    # Enable detailed HTTP request logging

# Authors filter: If set to a list of names, only process those authors.
# If empty list or None, process all QuSoft authors from the database.
# Examples:
# AUTHORS = ["Llorenç Escolà-Farràs", "Christian Schaffner"]  # Only these authors
AUTHORS = []  # Process all authors

# ============================================================================
# SETUP AND UTILITIES
# ============================================================================

# Create necessary directories
for directory in [CACHE_DIR, PDF_DOWNLOAD_DIR, TMP_DIR]:
    os.makedirs(directory, exist_ok=True)

# Configure logging
if ENABLE_HTTP_LOGGING:
    logging.basicConfig(level=logging.DEBUG)
    urllib3_logger = logging.getLogger('urllib3.connectionpool')
    urllib3_logger.setLevel(logging.DEBUG)

def normalize_name(name):
    """
    Normalize a name for comparison by converting to ASCII lowercase.
    Handles Unicode characters by removing diacritics.
    
    Args:
        name (str): The name to normalize
        
    Returns:
        str: Normalized name in ASCII lowercase
    """
    if not name:
        return ""
    return unicodedata.normalize('NFKD', name).encode('ascii', 'ignore').decode('ascii').lower()

def load_cache():
    """
    Load the query cache from disk.
    
    Returns:
        dict: Cache data, empty dict if file doesn't exist
    """
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding='utf-8') as f:
                return json.load(f)
        except (json.JSONDecodeError, FileNotFoundError):
            print(f"Warning: Could not load cache file {CACHE_FILE}, starting with empty cache")
            return {}
    return {}

def save_cache(cache):
    """
    Save the query cache to disk.
    
    Args:
        cache (dict): Cache data to save
    """
    try:
        with open(CACHE_FILE, "w", encoding='utf-8') as f:
            json.dump(cache, f, indent=2, ensure_ascii=False)
    except Exception as e:
        print(f"Warning: Could not save cache: {e}")

# ============================================================================
# QUSOFT DATABASE HANDLING
# ============================================================================

def get_qusoft_people_roles(file_path):
    """
    Load QuSoft people and their roles from the Excel database.
    
    Args:
        file_path (str): Path to the QuSoft people Excel file
        
    Returns:
        dict: Mapping of original names to roles
    """
    try:
        workbook = load_workbook(file_path)
        sheet = workbook.active
        people_roles = {}
        
        # Name in column 1, Role in column 5
        for row in sheet.iter_rows(min_row=3):
            name_cell = row[0]
            role_cell = row[4] if len(row) > 4 else None
            
            if name_cell.value:
                original_name = str(name_cell.value).strip()
                role = str(role_cell.value).strip() if role_cell and role_cell.value else "Unknown"
                people_roles[original_name] = role
                
        print(f"Loaded {len(people_roles)} QuSoft people from {file_path}")
        return people_roles
        
    except Exception as e:
        print(f"Error loading QuSoft people database: {e}")
        return {}

# ============================================================================
# ARXIV SEARCH FUNCTIONALITY
# ============================================================================

def generate_search_queries(person):
    """
    Generate search queries for a person using intelligent fallback strategies.
    
    This function implements a multi-strategy approach:
    1. Try standard author searches first (for exact matches)
    2. Fall back to AND combinations of distinctive name parts (for Unicode issues)
    
    Args:
        person (str): The person's name to search for
        
    Returns:
        list: List of search query strings to try in order
    """
    queries = []
    
    # Strategy 1: Standard author field searches
    queries.extend([
        f'au:"{person}"',  # Full name with quotes for exact match
        f'au:{person}',    # Full name without quotes for looser match
    ])
    
    # Strategy 2: AND combinations for Unicode-robust searching
    name_parts = person.replace('-', ' ').split()
    if len(name_parts) >= 2:
        # Convert to ASCII for better search compatibility
        ascii_parts = [normalize_name(part).title() for part in name_parts if len(part) > 3]
        
        if len(ascii_parts) >= 2:
            # Create combinations of the most distinctive parts
            first_part = ascii_parts[0]
            for other_part in ascii_parts[1:]:
                if len(other_part) >= 4:  # Use significant parts only
                    # Take first 5 chars for distinctiveness while avoiding over-specificity
                    partial_other = other_part[:5]
                    queries.append(f'{first_part} AND {partial_other}')
                else:
                    queries.append(f'{first_part} AND {other_part}')
    
    return queries

def verify_author_match(target_person, result_authors):
    """
    Verify if the target person is actually among the result authors.
    
    Uses the same normalization logic as the QuSoft author highlighting to ensure
    consistency. This prevents false positives from overly broad search queries.
    
    Args:
        target_person (str): The person we're looking for
        result_authors (list): List of arxiv.Author objects from a paper
        
    Returns:
        tuple: (is_match: bool, matching_author_name: str or None)
    """
    target_normalized = normalize_name(target_person)
    
    for author in result_authors:
        author_normalized = normalize_name(author.name)
        # Use exact normalized comparison (same as ** highlighting logic)
        if author_normalized == target_normalized:
            return True, author.name
    
    return False, None

def search_arxiv_for_authors(qusoft_people, cache):
    """
    Search arXiv for papers by QuSoft authors using intelligent search strategies.
    
    This function implements a sophisticated search approach:
    1. Uses multiple query strategies per author (exact, partial, AND combinations)
    2. Verifies that the target author is actually present in each result
    3. Highlights QuSoft authors with ** in the results
    4. Caches all results to avoid redundant API calls
    
    Args:
        qusoft_people (list): List of QuSoft people names to search for
        cache (dict): Cache dictionary for storing results
        
    Returns:
        list: List of paper dictionaries with metadata
    """
    client = arxiv.Client()
    results = []
    
    print(f"\\nSearching arXiv for papers by {len(qusoft_people)} QuSoft authors...")
    
    for i, person in enumerate(qusoft_people, 1):
        print(f"\\n[{i}/{len(qusoft_people)}] Processing: {person}")
        
        cache_key = f"arxiv:{normalize_name(person)}"
        if cache_key in cache:
            person_results = cache[cache_key]
            print(f"  ✓ Found {len(person_results)} cached papers")
        else:
            queries_to_try = generate_search_queries(person)
            person_results = []
            
            # Try each query strategy until we find results
            for query in queries_to_try:
                print(f"  Trying query: {query}")
                try:
                    search = arxiv.Search(
                        query=query,
                        max_results=25,  # Reasonable limit for partial searches
                        sort_by=arxiv.SortCriterion.SubmittedDate
                    )
                    
                    query_results = []
                    for result in client.results(search):
                        # Verify the target author is actually present (prevents false positives)
                        is_match, matching_author = verify_author_match(person, result.authors)
                        
                        if is_match:
                            # Create highlighted author list (QuSoft authors get **)
                            highlighted_authors = []
                            qusoft_author_count = 0
                            
                            for author in result.authors:
                                normalized_author = normalize_name(author.name)
                                # Check if this author is a QuSoft person
                                if normalized_author in [normalize_name(p) for p in qusoft_people]:
                                    highlighted_authors.append(f"**{author.name}**")
                                    qusoft_author_count += 1
                                else:
                                    highlighted_authors.append(author.name)
                            
                            # Create paper entry
                            entry = {
                                "source": "arXiv",
                                "id": result.entry_id.split("/")[-1],
                                "link": result.entry_id,
                                "authors": ", ".join(highlighted_authors),
                                "title": result.title,
                                "revision": str(result.updated.replace(tzinfo=None)),
                                "qusoft_author_count": qusoft_author_count
                            }
                            query_results.append(entry)
                    
                    if query_results:
                        person_results = query_results
                        print(f"  ✓ Found {len(person_results)} verified papers using: {query}")
                        break
                        
                except Exception as e:
                    print(f"  ✗ Error with query '{query}': {e}")
                    
            if not person_results:
                print(f"  \\033[91m✗ No papers found for {person}\\033[0m")
                
            # Cache the results (even if empty)
            cache[cache_key] = person_results
            save_cache(cache)
            
        results.extend(person_results)
    
    return results

# ============================================================================
# EXCEL OUTPUT FUNCTIONALITY
# ============================================================================

def save_to_excel(papers, output_file):
    """
    Save paper data to an Excel file with formatting and hyperlinks.
    
    Args:
        papers (list): List of paper dictionaries
        output_file (str): Output Excel file path
    """
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "QuSoft Papers"
        
        # Define headers
        headers = ["Source", "ID", "Authors", "Title", "Revision", "QuSoft Author Count"]
        sheet.append(headers)
        
        # Add data rows
        for entry in papers:
            row = [
                entry["source"],
                entry["id"],
                entry["authors"],
                entry["title"],
                entry["revision"],
                entry["qusoft_author_count"]
            ]
            sheet.append(row)
            
            # Make ID column clickable (hyperlink to paper)
            id_cell = sheet.cell(row=sheet.max_row, column=2)
            id_cell.hyperlink = entry["link"]
            id_cell.style = "Hyperlink"
        
        # Create formatted table
        if sheet.max_row > 1:  # Only if we have data
            table_range = f"A1:F{sheet.max_row}"
            table = Table(displayName="QuSoftPapersTable", ref=table_range)
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=True
            )
            table.tableStyleInfo = style
            sheet.add_table(table)
        
        workbook.save(output_file)
        print(f"\\n✓ Excel file saved: {output_file} ({len(papers)} papers)")
        
    except Exception as e:
        print(f"✗ Error saving Excel file: {e}")

# ============================================================================
# PDF GENERATION FUNCTIONALITY
# ============================================================================

def find_portrait(author):
    """
    Find a portrait image for an author.
    
    Args:
        author (str): Author name
        
    Returns:
        str or None: Path to portrait image, or None if not found
    """
    if not os.path.exists(PORTRAITS_DIR):
        return None
        
    key = normalize_name(author).replace(" ", "_").replace(".", "")
    
    try:
        for filename in os.listdir(PORTRAITS_DIR):
            name_part = os.path.splitext(filename)[0].lower()
            if key in name_part:
                return os.path.join(PORTRAITS_DIR, filename)
    except Exception as e:
        print(f"Warning: Error searching for portrait of '{author}': {e}")
    
    return None

def download_pdf(arxiv_id, cache):
    """
    Download a PDF from arXiv with caching.
    
    Args:
        arxiv_id (str): arXiv paper ID
        cache (dict): Cache dictionary
        
    Returns:
        str: Path to downloaded PDF file
    """
    pdf_url = f"https://arxiv.org/pdf/{arxiv_id}.pdf"
    pdf_path = os.path.join(PDF_DOWNLOAD_DIR, f"{arxiv_id}.pdf")
    cache_key = f"pdf:{arxiv_id}"
    
    if cache.get(cache_key) == "downloaded" and os.path.exists(pdf_path):
        return pdf_path
    
    try:
        print(f"  Downloading PDF: {arxiv_id}")
        response = requests.get(pdf_url, timeout=30)
        response.raise_for_status()
        
        with open(pdf_path, "wb") as f:
            f.write(response.content)
            
        cache[cache_key] = "downloaded"
        save_cache(cache)
        return pdf_path
        
    except Exception as e:
        print(f"  ✗ Error downloading PDF {arxiv_id}: {e}")
        return None

def create_combined_page(portrait_info, first_page_pdf_path):
    """
    Create a combined PDF page with portraits, highlights, and QR code.
    
    Args:
        portrait_info (list): List of (portrait_path, role, normalized_name) tuples
        first_page_pdf_path (str): Path to the original PDF first page
        
    Returns:
        fitz.Document: New document with the combined page
    """
    import fitz
    from PIL import Image
    
    # Open the first page of the PDF
    src_doc = fitz.open(first_page_pdf_path)
    src_page = src_doc[0]

    # Create a new blank page with same size
    width, height = src_page.rect.width, src_page.rect.height
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    
    # Scale down the first page and insert it lower (leave space for portraits)
    scale = 0.85
    new_width = width * scale
    new_height = height * scale
    x_offset = (width - new_width) / 2
    y_offset = 160  # Space for portraits at the top
    page.show_pdf_page(
        fitz.Rect(x_offset, y_offset, x_offset + new_width, y_offset + new_height),
        src_doc,
        0
    )

    # Highlight QuSoft-related words in yellow
    highlight_words = ["QuSoft", "CWI", "University of Amsterdam", "Centrum voor Wiskunde en Informatica"]
    for word in highlight_words:
        for inst in page.search_for(word, quads=True):
            highlight = page.add_highlight_annot(inst)
            highlight.set_colors(stroke=(1, 1, 0))  # Yellow
            highlight.update()

    # Highlight author names in light orange
    highlight_words = [norm_name for _, _, norm_name in portrait_info]
    for word in highlight_words:
        for inst in page.search_for(word, quads=True):
            highlight = page.add_highlight_annot(inst)
            highlight.set_colors(stroke=(1, 0.7, 0))  # Light orange
            highlight.update()

    # Check if QuSoft is mentioned on the first page
    first_page_text = src_page.get_text()
    if "qusoft" not in first_page_text.lower():
        print(f"    Warning: 'QuSoft' not found on first page")
        # Add warning annotation
        x_pos, y_pos = 10, height - 20
        rect = fitz.Rect(x_pos, y_pos, x_pos + 400, y_pos + 20)
        page.add_freetext_annot(rect, "'QuSoft' not mentioned on first page", 
                              fontsize=8, text_color=(1, 0, 0))

    # Add portraits as draggable stamps at the top
    x, y = 50, 40
    for portrait_path, role, _ in portrait_info:
        if not os.path.exists(portrait_path):
            print(f"    Warning: Portrait not found: {portrait_path}")
            continue
            
        # Set portrait size based on role
        role_str = str(role).lower() if role else ""
        if "students" in role_str or "support" in role_str:
            w, h = 120, 120
        elif "postdoc" in role_str:
            w, h = 100, 100
        else:
            w, h = 80, 80
            
        rect = fitz.Rect(x, y, x + w, y + h)
        
        # Convert image to RGB PNG for compatibility
        try:
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_img_png:
                img = Image.open(portrait_path).convert("RGB")
                img.save(tmp_img_png.name, "PNG")
                pix = fitz.Pixmap(tmp_img_png.name)
                page.add_stamp_annot(rect, stamp=pix)
                os.unlink(tmp_img_png.name)  # Clean up temp file
        except Exception as e:
            print(f"    Warning: Could not add portrait {portrait_path}: {e}")
            
        x += w + 20

    # Add QIP submission checkbox if enabled
    if SUBMISSION_TO_QIP_BOX:
        if 580 - x > 180:
            x_pos, y_pos = x, 70
        else:
            x_pos, y_pos = 300, 170
            
        page.add_freetext_annot(
            (x_pos, y_pos, x_pos + 160, y_pos + 40), 
            "submission to QIP2026?  YES             NO", 
            fontsize=12, text_color=(0, 0, 0)
        )
        page.draw_rect((x_pos + 33, y_pos + 15, x_pos + 43, y_pos + 25))
        page.draw_rect((x_pos + 91, y_pos + 15, x_pos + 101, y_pos + 25))

    # Add QR code with arXiv URL in lower-right corner
    try:
        arxiv_id = os.path.basename(first_page_pdf_path).replace('.pdf', '')
        qr_url = f"https://arxiv.org/abs/{arxiv_id}"
        qr_img = qrcode.make(qr_url)
        qr_size = 70
        
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_qr_png:
            qr_img = qr_img.resize((qr_size, qr_size))
            qr_img.save(tmp_qr_png.name, "PNG")
            pix = fitz.Pixmap(tmp_qr_png.name)
            
            # Place QR code in lower-right corner
            rect = fitz.Rect(width - qr_size - 15, height - qr_size - 20, 
                           width - 15, height - 20)
            page.add_stamp_annot(rect, stamp=pix)
            os.unlink(tmp_qr_png.name)  # Clean up temp file
            
    except Exception as e:
        print(f"    Warning: Could not add QR code: {e}")

    return doc

def generate_pdf_firstpages(papers, people_roles, cache):
    """
    Generate a PDF with annotated first pages for all papers.
    
    Args:
        papers (list): List of paper dictionaries
        people_roles (dict): Mapping of names to roles
        cache (dict): Cache dictionary
        
    Returns:
        str or None: Path to generated PDF, or None if failed
    """
    if not papers:
        print("\\nNo papers to generate PDF for")
        return None
    
    try:
        import fitz
    except ImportError:
        print("\\n✗ PyMuPDF not available, skipping PDF generation")
        return None
    
    print(f"\\nGenerating PDF first pages for {len(papers)} papers...")
    
    output_doc = fitz.open()
    
    for i, paper in enumerate(papers, 1):
        arxiv_id = paper["id"]
        print(f"\\n[{i}/{len(papers)}] Processing: {arxiv_id}")
        
        # Download PDF
        pdf_path = download_pdf(arxiv_id, cache)
        if not pdf_path:
            continue
        
        # Extract author names (remove ** highlighting)
        author_names = [a.strip().replace("**", "") for a in paper["authors"].split(",")]
        
        # Gather portrait information
        portrait_info = []
        for author_name in author_names:
            norm_author = normalize_name(author_name)
            portrait_path = find_portrait(author_name)
            
            # Find role by matching normalized names
            role = "Unknown"
            for original_name, person_role in people_roles.items():
                if normalize_name(original_name) == norm_author:
                    role = person_role
                    break
            
            if portrait_path:
                portrait_info.append((portrait_path, role, norm_author))
            else:
                print(f"    Warning: No portrait found for '{author_name}'")
        
        # Create combined page
        try:
            combined_doc = create_combined_page(portrait_info, pdf_path)
            output_doc.insert_pdf(combined_doc)
            print(f"    ✓ Added page with {len(portrait_info)} portraits")
        except Exception as e:
            print(f"    ✗ Error creating combined page: {e}")
            # Fall back to original PDF first page
            try:
                original_doc = fitz.open(pdf_path)
                output_doc.insert_pdf(original_doc, from_page=0, to_page=0)
                print(f"    ✓ Added original first page as fallback")
            except Exception as e2:
                print(f"    ✗ Error adding fallback page: {e2}")
    
    # Save the final PDF
    if output_doc.page_count > 0:
        try:
            output_doc.save(PDF_OUTPUT)
            print(f"\\n✓ PDF saved: {PDF_OUTPUT} ({output_doc.page_count} pages)")
            return PDF_OUTPUT
        except Exception as e:
            print(f"\\n✗ Error saving PDF: {e}")
            return None
    else:
        print(f"\\n✗ No pages to save in PDF")
        return None

# ============================================================================
# MAIN FUNCTIONALITY
# ============================================================================

def filter_papers_by_year(papers, starts_with):
    """
    Filter papers by year prefix.
    
    Args:
        papers (list): List of paper dictionaries
        starts_with (str): Year prefix to filter by
        
    Returns:
        list: Filtered papers
    """
    filtered = [p for p in papers if p["id"].startswith(starts_with)]
    
    # Remove duplicates by ID
    seen_ids = set()
    unique_papers = []
    for paper in filtered:
        if paper["id"] not in seen_ids:
            unique_papers.append(paper)
            seen_ids.add(paper["id"])
    
    return unique_papers

def main():
    """
    Main function that orchestrates the entire process.
    """
    print("=" * 70)
    print("QuSoft arXiv Paper Scanner and PDF Generator")
    print("=" * 70)
    
    # Load QuSoft people database
    if not os.path.exists(QUSOFT_PEOPLE_FILE):
        print(f"✗ Error: QuSoft people database not found: {QUSOFT_PEOPLE_FILE}")
        return
    
    people_roles = get_qusoft_people_roles(QUSOFT_PEOPLE_FILE)
    if not people_roles:
        print("✗ Error: No people loaded from database")
        return
    
    all_qusoft_people = list(people_roles.keys())
    
    # Apply author filter if specified
    if AUTHORS and len(AUTHORS) > 0:
        qusoft_people = [name for name in AUTHORS if name in all_qusoft_people]
        missing_authors = [name for name in AUTHORS if name not in all_qusoft_people]
        
        print(f"\\nAUTHORS filter active: Processing {len(qusoft_people)} out of {len(all_qusoft_people)} QuSoft authors")
        print(f"Selected authors: {qusoft_people}")
        
        if missing_authors:
            print(f"⚠️  Warning: These authors from AUTHORS list are not in the database: {missing_authors}")
    else:
        qusoft_people = all_qusoft_people
        print(f"\\nProcessing all {len(qusoft_people)} QuSoft authors from database")
    
    if not qusoft_people:
        print("✗ Error: No authors to process")
        return
    
    # Load cache and search arXiv
    cache = load_cache()
    arxiv_papers = search_arxiv_for_authors(qusoft_people, cache)
    
    print(f"\\n" + "=" * 50)
    print(f"RESULTS SUMMARY")
    print(f"=" * 50)
    print(f"Total papers found: {len(arxiv_papers)}")
    
    if arxiv_papers:
        # Show paper IDs for debugging
        print(f"\\nPaper IDs found:")
        for paper in arxiv_papers:
            print(f"  - {paper['id']} ({paper['title'][:50]}...)")
        
        # Filter by year if specified
        if STARTS_WITH:
            print(f"\\nFiltering for papers starting with '{STARTS_WITH}'...")
            filtered_papers = filter_papers_by_year(arxiv_papers, STARTS_WITH)
            print(f"Papers after year filter: {len(filtered_papers)}")
        else:
            filtered_papers = arxiv_papers
        
        # Generate Excel output
        save_to_excel(arxiv_papers, EXCEL_OUTPUT)  # All papers
        
        # Generate PDF output for filtered papers
        if filtered_papers:
            # Sort by ID (newest first)
            filtered_papers.sort(key=lambda p: p["id"], reverse=True)
            generate_pdf_firstpages(filtered_papers, people_roles, cache)
        else:
            print(f"\\n⚠️  No papers match the year filter '{STARTS_WITH}'")
            print(f"   Consider adjusting STARTS_WITH or removing the filter")
    else:
        print("\\nNo papers found for any QuSoft authors")
    
    print(f"\\n" + "=" * 70)
    print("Process completed!")
    print(f"=" * 70)

if __name__ == "__main__":
    main()