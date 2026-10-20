#!/usr/bin/env python3
"""Export papers to an Excel sheet (for reports), with clickable arXiv links.

    python scripts/export_excel.py                       # all papers
    python scripts/export_excel.py --year 2025
    python scripts/export_excel.py --start-date 2025-09-01 --end-date 2025-10-31
    python scripts/export_excel.py --authors "Christian Schaffner"
"""

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from qusoft.common import filter_papers, is_alumnus, load_papers, load_people, path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--year", type=int)
    ap.add_argument("--month", help="YYYY-MM")
    ap.add_argument("--start-date", help="YYYY-MM-DD")
    ap.add_argument("--end-date", help="YYYY-MM-DD")
    ap.add_argument("--authors", nargs="+", help="only papers with these QuSoft authors")
    ap.add_argument("--output", help="default: output/qusoft_papers[_filter].xlsx")
    args = ap.parse_args()

    people = load_people()
    papers = filter_papers(load_papers(), args.year, args.month, args.start_date, args.end_date, args.authors)
    suffix = f"_{args.year}" if args.year else f"_{args.month}" if args.month else ""
    out = Path(args.output) if args.output else path("paths", "output_dir") / f"qusoft_papers{suffix}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "QuSoft papers"
    ws.append(["arXiv id", "Submitted", "Title", "Authors", "QuSoft authors", "Of whom alumni", "Subjects"])
    for p in papers:
        q = p["qusoft_authors"]
        ws.append([
            p["id"],
            p["submitted"][:10],
            p["title"],
            ", ".join(a["name"] for a in p["authors"]),
            ", ".join(q),
            ", ".join(n for n in q if is_alumnus(people.get(n, {}))),
            ", ".join(p.get("categories", [])),
        ])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.hyperlink = p["link"]
        cell.font = Font(color="0563C1", underline="single")

    for col, width in zip("ABCDEFG", (14, 12, 70, 60, 40, 25, 20)):
        ws.column_dimensions[col].width = width
    if papers:
        table = Table(displayName="QuSoftPapers", ref=f"A1:G{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"Saved {len(papers)} papers to {out}")


if __name__ == "__main__":
    main()
