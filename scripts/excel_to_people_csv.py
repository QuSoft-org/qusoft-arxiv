#!/usr/bin/env python3
"""Refresh people.csv from the private QuSoft Database.xlsx.

Only people marked "Featured on Page = Yes" are exported, and only the fields
that are fine to publish: name, role, active/alumni, start year and end year.
E-mail addresses, supervisors and exact dates stay in the spreadsheet.

    python scripts/excel_to_people_csv.py
"""

import csv
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from qusoft.common import PEOPLE_FIELDS, path

# Columns in the spreadsheet (0-based). Row 1 is a warning banner, row 2 the header.
NAME, FEATURED, ROLE, STATUS, START, END = 0, 2, 4, 5, 10, 11


def year_of(value):
    """Year from a date cell, which may be a datetime, a number like 2025.0, or text."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, (int, float)) and 1990 < value < 2100:
        return int(value)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y"):
        try:
            return datetime.strptime(text, fmt).year
        except ValueError:
            pass
    return None


def main():
    excel = path("database", "excel_file")
    sheet = load_workbook(excel, read_only=True).active
    rows, skipped = [], []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        name = " ".join(str(row[NAME] or "").split())
        if not name or str(row[FEATURED] or "").strip() != "Yes":
            continue
        if "://" in name or any(ch.isdigit() for ch in name):
            skipped.append(name)
            continue
        rows.append({
            "name": name,
            "role": str(row[ROLE] or "Unknown").strip(),
            "status": "alumni" if str(row[STATUS] or "").strip().lower() == "alumni" else "active",
            "start_year": year_of(row[START]) or "",
            "end_year": year_of(row[END]) or "",
        })

    rows.sort(key=lambda r: r["name"])
    out = path("database", "people_csv")
    with open(out, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=PEOPLE_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    alumni = sum(r["status"] == "alumni" for r in rows)
    print(f"Wrote {len(rows)} people ({alumni} alumni) to {out}")
    for name in skipped:
        print(f"Skipped a row whose name doesn't look like a name: {name!r}. Fix it in the spreadsheet.")


if __name__ == "__main__":
    main()
