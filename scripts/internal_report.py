#!/usr/bin/env python3
"""Internal report for the QuSoft management team (local only, never published).

Looks at every arXiv paper by a QuSoft person written while they were at QuSoft,
including the ones that don't name QuSoft, and counts papers
per person per year. Membership dates come from the private spreadsheet when it
is present (exact dates), otherwise from the years in people.csv.

    python scripts/internal_report.py            # -> output/internal_report_<date>.xlsx

It needs the papers the daily search rejected (data/rejected_ids.json). Their
details are fetched once from arXiv and kept in cache/, which is not in git.
"""

import json
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import arxiv
import fitz
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from qusoft.firstpage import qusoft_page
from qusoft.common import (PersonMatcher, base_id, is_alumnus, load_json, load_papers, load_people,
                           path, save_json)

ROOT = Path(__file__).resolve().parent.parent
META_CACHE = ROOT / "cache" / "report_papers.json"
OLD_QUERY_CACHE = ROOT / "cache" / "query_cache.json"
UA = {"User-Agent": "qusoft-arxiv internal report (https://github.com/cschaffner/qusoft-arxiv)"}

QUSOFT_ON_FIRST_PAGE = "QuSoft named (on the site)"  # name kept from when only the first page counted
QUSOFT_LATER = "QuSoft only thanked"
OLD_KINDS = {"QuSoft on first page": QUSOFT_ON_FIRST_PAGE, "QuSoft only on a later page": QUSOFT_LATER}
CWI_UVA = "CWI or UvA, no QuSoft"
NONE = "No QuSoft, CWI or UvA"
UNREADABLE = "PDF not available"
KINDS = [QUSOFT_ON_FIRST_PAGE, QUSOFT_LATER, CWI_UVA, NONE, UNREADABLE]
CWI_UVA_WORDS = re.compile(r"\bCWI\b|centrum wiskunde|university of amsterdam|universiteit van amsterdam|\bUvA\b", re.I)


# ---------------------------------------------------------------------------
# Membership dates
# ---------------------------------------------------------------------------

def _date(value, end=False):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, (int, float)) and 1990 < value < 2100:
        return date(int(value), 12, 31) if end else date(int(value), 1, 1)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    if re.fullmatch(r"\d{4}", text):
        return date(int(text), 12, 31) if end else date(int(text), 1, 1)
    return None


def membership(people):
    """{name: (start, end)}; end None means still at QuSoft, start None means unknown."""
    spans = {n: (_date(i.get("start_year")), _date(i.get("end_year"), end=True)) for n, i in people.items()}
    excel = path("database", "excel_file")
    source = "years from people.csv"
    if excel.exists():
        sheet = load_workbook(excel, read_only=True).active
        for row in sheet.iter_rows(min_row=3, values_only=True):  # same columns as excel_to_people_csv.py
            name = " ".join(str(row[0] or "").split())
            if name in spans:
                start, end = _date(row[10]), _date(row[11], end=True)
                spans[name] = (start or spans[name][0], end or spans[name][1])
        source = "exact dates from the QuSoft spreadsheet"
    return spans, source


def was_member(span, day, alumnus):
    start, end = span
    if start and day < start:
        return False
    # Current members can be past their planned end date (e.g. a PhD running late).
    return not (end and day > end and alumnus)


# ---------------------------------------------------------------------------
# Rejected papers: details and affiliation check
# ---------------------------------------------------------------------------

def _from_old_cache():
    if not OLD_QUERY_CACHE.exists():
        return {}
    found = {}
    for entries in json.loads(OLD_QUERY_CACHE.read_text()).values():
        for e in entries if isinstance(entries, list) else []:
            found[base_id(e["id"])] = {
                "id": e["id"], "title": e["title"], "submitted": e["submitted"][:10],
                "authors": [a.strip(" *") for a in e["authors"].split(",") if a.strip(" *")],
                "categories": e.get("categories", []),
            }
    return found


def _fetch_details(ids):
    client = arxiv.Client(page_size=50, delay_seconds=5, num_retries=6)
    out = {}
    for i in range(0, len(ids), 50):
        batch = ids[i:i + 50]
        for attempt in range(5):
            try:
                for r in client.results(arxiv.Search(id_list=batch)):
                    out[base_id(r.get_short_id())] = {
                        "id": r.get_short_id(), "title": r.title, "submitted": r.published.date().isoformat(),
                        "authors": [a.name for a in r.authors],
                        "categories": list(r.categories),
                    }
                break
            except arxiv.HTTPError as e:
                wait = 30 * (attempt + 1)
                print(f"    arXiv said {e.status}, waiting {wait}s")
                time.sleep(wait)
    return out


def _pdf(arxiv_id):
    cache = path("paths", "local_pdf_cache")
    cache.mkdir(parents=True, exist_ok=True)
    for f in [cache / f"{arxiv_id}.pdf", *cache.glob(f"{base_id(arxiv_id)}v*.pdf")]:
        if f.exists():
            return f
    target = cache / f"{arxiv_id}.pdf"
    try:
        r = requests.get(f"https://arxiv.org/pdf/{arxiv_id}", headers=UA, timeout=60)
        time.sleep(3)
        if r.ok and r.content[:4] == b"%PDF":
            target.write_bytes(r.content)
            return target
    except requests.RequestException:
        pass
    return None


def affiliation_kind(pdf):
    try:
        with fitz.open(pdf) as doc:
            first = doc[0].get_text()
            if qusoft_page(doc):
                return QUSOFT_ON_FIRST_PAGE
            if any("qusoft" in page.get_text().lower() for page in doc):
                return QUSOFT_LATER
            return CWI_UVA if CWI_UVA_WORDS.search(first) else NONE
    except Exception:
        return UNREADABLE


def rejected_papers():
    ids = [base_id(i) for i in load_json(path("database", "rejected_ids"), [])]
    cache = load_json(META_CACHE, {})
    todo = [i for i in ids if "categories" not in cache.get(i, {})]
    if todo:
        old = _from_old_cache()
        for i in todo:
            if i in old:
                cache[i] = {**old[i], **{k: v for k, v in cache.get(i, {}).items() if k == "kind"}}
        todo = [i for i in ids if "categories" not in cache.get(i, {})]
        if todo:
            print(f"Fetching details of {len(todo)} papers from arXiv")
            for i, details in _fetch_details(todo).items():
                cache[i] = {**cache.get(i, {}), **details}
        save_json(META_CACHE, cache)
    missing_kind = [i for i in ids if i in cache and "kind" not in cache[i]]
    if missing_kind:
        print(f"Checking affiliations in {len(missing_kind)} PDFs")
        for n, i in enumerate(missing_kind, 1):
            pdf = _pdf(cache[i]["id"])
            cache[i]["kind"] = affiliation_kind(pdf) if pdf else UNREADABLE
            if n % 50 == 0:
                print(f"  {n}/{len(missing_kind)}")
                save_json(META_CACHE, cache)
        save_json(META_CACHE, cache)
    return [cache[i] for i in ids if i in cache]


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

def collect(people):
    spans, source = membership(people)
    match = PersonMatcher(people)
    rows = []  # one per paper: date, id, title, kind, members (QuSoft authors who were members then)

    for p in load_papers():
        d = date.fromisoformat(p["submitted"][:10])
        members = [a for a in dict.fromkeys(p["qusoft_authors"]) if a in people]
        rows.append({"date": d, "id": p["id"], "title": p["title"], "kind": QUSOFT_ON_FIRST_PAGE,
                     "members": members, "authors": [a["name"] for a in p["authors"]],
                     "categories": p.get("categories", [])})

    for p in rejected_papers():
        d = date.fromisoformat(p["submitted"][:10])
        members = [m for m in dict.fromkeys(filter(None, map(match, p["authors"])))
                   if was_member(spans[m], d, is_alumnus(people[m]))]
        if members:
            kind = OLD_KINDS.get(p.get("kind"), p.get("kind", UNREADABLE))
            if kind == QUSOFT_ON_FIRST_PAGE:  # QuSoft is there after all (e.g. a later version)
                kind = QUSOFT_LATER
            rows.append({"date": d, "id": p["id"], "title": p["title"], "kind": kind,
                         "members": members, "authors": p["authors"], "categories": p.get("categories", [])})

    # A paper outside every subject someone has a QuSoft paper in may be by a namesake.
    subjects = defaultdict(set)
    for r in rows:
        if r["kind"] == QUSOFT_ON_FIRST_PAGE:
            for m in r["members"]:
                subjects[m].update(c.split(".")[0] for c in r["categories"])
    for r in rows:
        mine = {c.split(".")[0] for c in r["categories"]}
        r["unsure"] = [m for m in r["members"] if r["kind"] != QUSOFT_ON_FIRST_PAGE
                       and subjects[m] and mine and not mine & subjects[m]]
    return sorted(rows, key=lambda r: r["date"], reverse=True), source


def _sheet(wb, title, header, rows, widths, note=None):
    ws = wb.create_sheet(title)
    start = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="626C75")
        start = 3
    for c, h in enumerate(header, 1):
        cell = ws.cell(start, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C02127")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
                cell.data_type = "s"  # titles come from arXiv: never let Excel read them as formulas
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(start + 1, 2 if title.startswith("Per person") else 1)
    return ws


def main():
    people = load_people()
    rows, source = collect(people)
    years = sorted({r["date"].year for r in rows}, reverse=True)
    wb = Workbook()
    wb.remove(wb.active)

    # Per year: how often QuSoft is on the paper.
    per_year = Counter((r["date"].year, r["kind"]) for r in rows)
    summary = []
    for y in years:
        counts = [per_year[(y, k)] for k in KINDS]
        total = sum(counts)
        unsure = sum(1 for r in rows if r["date"].year == y and r["unsure"] and set(r["unsure"]) == set(r["members"]))
        summary.append([y, total, *counts, round(100 * counts[0] / total) if total else None, unsure])
    _sheet(wb, "Per year", ["Year", "Papers by members", *KINDS, "% naming QuSoft",
                            "Of these, maybe by someone else with the same name"], summary,
           [8, 12, 14, 14, 14, 14, 12, 14, 18],
           note=f"Papers by QuSoft people written while they were members ({source}). "
                "Only papers that QuSoft's daily arXiv search found.")

    # The papers that don't name QuSoft.
    without = [[r["date"], f"https://arxiv.org/abs/{base_id(r['id'])}", r["title"], r["kind"],
                ", ".join(r["members"]), ", ".join(r["unsure"]) or None, ", ".join(r["categories"]),
                ", ".join(r["authors"])]
               for r in rows if r["kind"] != QUSOFT_ON_FIRST_PAGE]
    ws = _sheet(wb, "Without QuSoft",
                ["Date", "arXiv", "Title", "Affiliation", "QuSoft members", "Maybe someone else with this name",
                 "arXiv subjects", "All authors"],
                without, [11, 32, 60, 26, 30, 22, 22, 60],
                note="A different researcher with the same name can end up here. The 'maybe someone else' column "
                     "names members whose QuSoft papers are all in other arXiv subjects; check those by eye.")
    for row in ws.iter_rows(min_row=4, max_col=1):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    for row in ws.iter_rows(min_row=4, min_col=2, max_col=2):
        for cell in row:
            cell.hyperlink, cell.font = cell.value, Font(color="921317", underline="single")

    # Per person per year.
    total, with_qusoft = defaultdict(Counter), defaultdict(Counter)
    for r in rows:
        for m in r["members"]:
            if m in r["unsure"]:
                continue
            total[m][r["date"].year] += 1
            if r["kind"] == QUSOFT_ON_FIRST_PAGE:
                with_qusoft[m][r["date"].year] += 1
    names = sorted(people, key=lambda n: (is_alumnus(people[n]), n))
    for title, counts, note in [
        ("Per person, all", total, "Papers per person per year while a member, with or without QuSoft as affiliation. Papers that may be by someone else with the same name are left out."),
        ("Per person, QuSoft", with_qusoft, "Papers per person per year that name QuSoft."),
    ]:
        table = [[n, people[n].get("role", ""), "alumni" if is_alumnus(people[n]) else "current",
                  sum(counts[n].values()), *[counts[n][y] or None for y in years]] for n in names]
        _sheet(wb, title, ["Name", "Role", "Status", "Total", *years], table,
               [28, 18, 9, 7, *[6] * len(years)], note=note)

    out = path("paths", "output_dir") / f"internal_report_{date.today().isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")
    for y, t, *c, pct, _ in summary[:6]:
        print(f"  {y}: {t} papers by members, {c[0]} naming QuSoft ({pct}%)")


if __name__ == "__main__":
    main()
